"""Extract mixed local files with stable human-readable locators; never execute source content."""
import sys,json,hashlib,csv
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

def extract(folder):
 sources=[]
 for p in sorted(Path(folder).iterdir()):
  if not p.is_file(): continue
  obj={'id':p.stem.split('_')[0], 'path':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'status':'extracted','items':[]}
  def add(loc,text,**extra):obj['items'].append({'locator':loc,'text':str(text),**extra})
  try:
   ext=p.suffix.lower()
   if ext=='.docx':
    d=Document(p)
    for i,para in enumerate([x for x in d.paragraphs if x.text.strip()],1):add(f'paragraph {i}',para.text)
    for ti,t in enumerate(d.tables,1):
     for ri,row in enumerate(t.rows,1):
      for ci,c in enumerate(row.cells,1):add(f'table {ti} row {ri} col {ci}',c.text)
   elif ext=='.xlsx':
    w=load_workbook(p,data_only=True); wf=load_workbook(p,data_only=False)
    for sh in w:
     for row in sh:
      for cell in row:
       f=wf[sh.title][cell.coordinate].value
       if cell.value is not None or f is not None:add(sh.title+'!'+cell.coordinate,cell.value,value=cell.value,formula=f if isinstance(f,str) and f.startswith('=') else None)
   elif ext in {'.txt','.md','.eml','.json','.csv','.tsv'}:
    for i,line in enumerate(p.read_text(encoding='utf-8-sig').splitlines(),1):
     if line.strip():add(f'line {i}',line)
   else:obj['status']='unsupported';obj['reason']='Use an appropriate format reader or OCR; this file was not reviewed.'
  except Exception as e:obj['status']='error';obj['reason']=str(e)
  sources.append(obj)
 return sources
if __name__=='__main__':
 if len(sys.argv)!=3:raise SystemExit('Usage: extract_sources.py INPUT_DIR OUTPUT_JSON')
 dest=Path(sys.argv[2]);dest.parent.mkdir(parents=True,exist_ok=True);result=extract(sys.argv[1]);dest.write_text(json.dumps({'sources':result},indent=2,ensure_ascii=False),encoding='utf-8');print(f'Extracted {sum(s["status"]=="extracted" for s in result)} / {len(result)} files -> {dest}')
