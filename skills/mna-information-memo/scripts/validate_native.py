"""Portable structure/content checks. Not a visual or desktop PowerPoint test."""
import json,sys,zipfile,re,hashlib,io,posixpath
from pathlib import Path
from xml.etree import ElementTree as E
from openpyxl import load_workbook
story=json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'));ppt=Path(sys.argv[2]);report=Path(sys.argv[3])
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
norm=lambda s:re.sub(r'\s+',' ',s).strip()
def words(node):return norm(' '.join(t.text or '' for t in node.findall('.//a:t',ns)))
counts={'slides':len(story['slides']),'native_tables':0,'native_charts':0,'native_text_elements':0,'source_notes':0,'embedded_chart_workbooks':0}
with zipfile.ZipFile(ppt) as z:
 assert z.testzip() is None,'Damaged PPTX ZIP'
 names=z.namelist();parts=[n for n in names if re.fullmatch(r'ppt/slides/slide\d+.xml',n)]
 assert len(parts)==len(story['slides']),'Slide count changed'
 for n,s in enumerate(story['slides'],1):
  page=E.fromstring(z.read(f'ppt/slides/slide{n}.xml'))
  assert not page.findall('.//p:pic',ns),'Unsupported raster replacement'
  tables=page.findall('.//a:tbl',ns);expected=[e for e in s['elements'] if e['type']=='table']
  assert len(tables)==len(expected),(n,'table count')
  for table,e in zip(tables,expected):
   rows=[[words(c) for c in r.findall('a:tc',ns)] for r in table.findall('a:tr',ns)]
   assert rows==[[norm(str(c)) for c in row] for row in e['rows']],(n,e['id'],'table content')
   counts['native_tables']+=1
  text=[words(sh) for sh in page.findall('.//p:sp',ns)]
  for e in s['elements']:
   if e['type']=='text':assert norm(e['text']) in text,(n,e['id'],'text content');counts['native_text_elements']+=1
  notes=words(E.fromstring(z.read(f'ppt/notesSlides/notesSlide{n}.xml')))
  for note in s.get('notes',[]):assert norm(note) in notes,(n,'missing source note')
  counts['source_notes']+=1
  assert len(page.findall('.//c:chart',ns))==sum(e['type']=='chart' for e in s['elements']),(n,'chart count')
 charts=sorted([n for n in names if re.fullmatch(r'ppt/charts/chart\d+.xml',n)],key=lambda n:int(re.search(r'(\d+)\.xml',n).group(1)))
 expected=[e for s in story['slides'] for e in s['elements'] if e['type']=='chart']
 assert len(charts)==len(expected)
 for name,e in zip(charts,expected):
  chart=E.fromstring(z.read(name));values=[float(v.text) for v in chart.findall('.//c:val/c:numRef/c:numCache/c:pt/c:v',ns)]
  labels=[v.text for v in chart.findall('.//c:cat//c:pt/c:v',ns)]
  assert values==e['values'] and labels==list(map(str,e['categories'])),(name,'chart data')
  external=chart.find('c:externalData',ns)
  assert external is not None,(name,'missing workbook relationship')
  relid=external.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
  rels=E.fromstring(z.read(posixpath.join(posixpath.dirname(name),'_rels',posixpath.basename(name)+'.rels')))
  target=next(r.attrib['Target'] for r in rels if r.attrib['Id']==relid)
  bookpath=posixpath.normpath(posixpath.join(posixpath.dirname(name),target))
  book=load_workbook(io.BytesIO(z.read(bookpath)),data_only=True)
  sheet=book.active;records=list(sheet.values)
  assert [str(r[0]) for r in records[1:]]==list(map(str,e['categories'])),(name,'workbook labels')
  assert [float(r[1]) for r in records[1:]]==e['values'],(name,'workbook values')
  counts['native_charts']+=1
 books=[n for n in names if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')]
 assert len(books)==len(expected),'Chart workbooks missing'
 for name in books:
  book=load_workbook(io.BytesIO(z.read(name)),data_only=True);assert book.sheetnames
 counts['embedded_chart_workbooks']=len(books)
result={'passed':counts,'sha256':hashlib.sha256(ppt.read_bytes()).hexdigest(),'renderer':'pptxgenjs','limitations':['Structural/content checks only. Render and inspect all pages separately.','Windows desktop PowerPoint edit/save/reopen has not been tested.','Charts contain editable snapshots, not live source-workbook connections.']}
report.write_text(json.dumps(result,indent=2),encoding='utf-8');print(json.dumps(result))
