from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import json, csv
root=Path(__file__).resolve().parents[1]; out=root/'inputs'
doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.add_heading('Northbridge Field Services company briefing',0)
doc.add_paragraph('FICTIONAL COMPANY AND DATA. Training simulation only. Prepared 15 January 2026. Preliminary management briefing, version 1.0.')
sections=[
('Company background','Northbridge Field Services Ltd was founded in 2012. The UK business provides equipment maintenance and technical field support to industrial and logistics customers. Its headquarters are in Birmingham. The transaction perimeter is the entire standalone UK company.'),
('Services and customers','The business supplies scheduled maintenance under annual service contracts, reactive repairs and installation projects. Contracts generally run for one to three years and renew by agreement. Scheduled maintenance provides the contracted revenue stream. No assumption should be made that renewals are automatic. Customers include logistics operators, manufacturers and facilities managers.'),
('Preliminary financial snapshot','The preliminary FY2025 revenue estimate is GBP 26.2 million. This briefing does not provide a reconciliation to the finance model. Financial years end on 31 December. All revenue figures exclude VAT.'),
('Operations and management','At 31 December 2025, Northbridge has 120 employees including 85 field engineers. It operates hubs in Birmingham, Manchester and Bristol. Maya Reed is founder and CEO, Daniel Cole is CFO and Aisha Patel is COO. These names and all company details are fictitious. The founder continues to lead major customer relationships.'),
('Growth plan','Management plans to add 20 net field engineers by the end of FY2027, cross-sell scheduled maintenance to repair customers and expand service coverage from existing hubs. The plan assumes sufficient technician recruitment and contract renewals. No signed pipeline or customer-cohort evidence accompanies this briefing.'),
('Ownership and proposed transaction','The founder holds 70%, management holds 20% and a seed investor holds 10% of ordinary equity. Shareholders are considering a sale of 70% to 100% of the company. Management rollover, valuation, timing and financing remain to be agreed. No debt, cash or working-capital schedule has been provided.'),
('Market positioning','Management describes the UK equipment-maintenance market as fragmented. No independent market-size study, market-share estimate or competitor benchmarking has been provided.')]
for title,body in sections:
 if title=='Growth plan': doc.add_page_break()
 doc.add_heading(title,1); doc.add_paragraph(body)
sec.footer.paragraphs[0].text='FICTIONAL TRAINING MATERIAL | Northbridge Field Services'

from docx.oxml.ns import qn
for node in [doc.element, doc.styles.element]:
 for border in list(node.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
doc.save(out/'01_company_brief.docx')
(out/'03_management_interview.txt').write_text('''FICTIONAL TRAINING MATERIAL\nNorthbridge management interview, 20 January 2026\n\n[00:03] CEO: The business was founded in 2012. Our headquarters are Birmingham. We serve UK industrial and logistics customers.\n[00:08] CFO: The FY2025 management P&L model shows revenue of GBP 25.0m and EBITDA of GBP 5.0m. The earlier company briefing shows GBP 26.2m. The difference may involve orders that were not recognised as revenue. We have not completed the reconciliation.\n[00:12] COO: At year end we had 120 employees, including 85 engineers, across Birmingham, Manchester and Bristol.\n[00:17] CEO: We estimate customer retention at 94% in FY2025, but have not supplied the cohort calculation or defined whether that refers to customers or revenue.\n[00:21] CEO: We believe the market is fragmented. We have not commissioned a market-sizing study and cannot substantiate a market-share percentage.\n[00:25] CFO: FY2026 and FY2027 are management forecasts, not contracted revenue. Revenue is expected to reach GBP 29.0m and GBP 33.6m, with EBITDA of GBP 6.1m and GBP 7.4m respectively.\n[00:28] CEO: Growth depends on technician recruitment, renewals and converting more repair customers to maintenance contracts. We plan 20 net new engineers by FY2027.\n[00:31] CFO: EBITDA is management-reported and unaudited. No normalisation bridge or quality-of-earnings report is available. We have not supplied cash, debt or normalised working capital.\n''')
(out/'04_cfo_email.eml').write_text('''From: cfo@northbridge.example\nTo: dealteam@example.invalid\nDate: Wed, 21 Jan 2026 09:00:00 +0000\nSubject: Fictional Northbridge financial model basis\nContent-Type: text/plain; charset=utf-8\n\nFICTIONAL TRAINING MATERIAL\nPlease use the attached management model as the provisional basis for the draft. The GBP 1.2m difference against the briefing remains unresolved and must remain visible in the draft. This is not confirmation that the model has been audited.\n\nThe P&L tab is in GBP millions. The CustomerSales and ServiceMix tabs are in GBP thousands. They cover FY2025 recognised revenue, excluding VAT, for the same UK company. CustomerSales contains the five largest customers plus one aggregate row for all others. Do not treat that aggregate as a single customer.\n\nThe model is version 2.0, dated 21 January 2026. The forecast is management's base case. There is no customer-level forecast bridge or recruitment cost schedule yet.\n''')
seed={'financials':[['Revenue',18,21.5,25,29,33.6],['Cost of services',-10.8,-12.7,-14.5,-16.5,-18.6],['Gross profit',None,None,None,None,None],['Operating expenses',-4.5,-4.9,-5.5,-6.4,-7.6],['EBITDA',None,None,None,None,None],['EBITDA margin',None,None,None,None,None]],'customers':[['C001','Customer A',3000],['C002','Customer B',2250],['C003','Customer C',2000],['C004','Customer D',1750],['C005','Customer E',1500],['OTHER','All other customers (aggregate)',14500]],'services':[['Scheduled maintenance',17500,'Annual service contracts'],['Reactive repairs',5000,'Call-out and materials'],['Installation projects',2500,'Project-based work']]}
(root/'.build/source_seed.json').write_text(json.dumps(seed,indent=2))
print('Created briefing, interview, email and financial seed')
